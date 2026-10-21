"""
Main application window.
"""
from __future__ import annotations
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from PyQt6.QtCore import Qt, QPropertyAnimation, QEasingCurve
from PyQt6.QtGui import QColor, QPalette, QLinearGradient, QGradient, QPainter, QFont, QIcon
from PyQt6.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QTabWidget, QProgressBar,
    QMessageBox, QFileDialog, QFrame, QSizePolicy,
    QScrollArea,
)

from core.models import CheckReport
from ui.check_worker import CheckWorker
from ui.widgets.file_upload_panel import FileUploadPanel
from ui.widgets.summary_bar import SummaryBar
from ui.widgets.result_table import RegionResultTable, DeviceResultTable


# ---------------------------------------------------------------------------
# Export colours (openpyxl PatternFill)
# ---------------------------------------------------------------------------
_EXPORT_FILLS: dict[str, PatternFill] = {
    "OK":                  PatternFill("solid", fgColor="0D2018"),
    "MISSING":             PatternFill("solid", fgColor="2A0F0F"),
    "UNSUPPORTED":         PatternFill("solid", fgColor="2A1F0A"),
    "NO_AVL_ID":           PatternFill("solid", fgColor="2A1500"),
    "POSSIBLE_MATCH":      PatternFill("solid", fgColor="251800"),
    "UNKNOWN_MAPPING":     PatternFill("solid", fgColor="0F1030"),
    "UNKNOWN_CATEGORY":    PatternFill("solid", fgColor="100F30"),
    "NOT_IN_AVL":          PatternFill("solid", fgColor="1A0A2A"),
    "AVL_NOT_QUALIFIED":   PatternFill("solid", fgColor="1A1500"),
}
_EXPORT_FONTS: dict[str, Font] = {
    "OK":                  Font(color="4ADE80", bold=True),
    "MISSING":             Font(color="F87171", bold=True),
    "UNSUPPORTED":         Font(color="FBBF24", bold=True),
    "NO_AVL_ID":           Font(color="F97316", bold=True),
    "POSSIBLE_MATCH":      Font(color="FB923C", bold=True),
    "UNKNOWN_MAPPING":     Font(color="818CF8", bold=True),
    "UNKNOWN_CATEGORY":    Font(color="A78BFA", bold=True),
    "NOT_IN_AVL":          Font(color="E879F9", bold=True),
    "AVL_NOT_QUALIFIED":   Font(color="FACC15", bold=True),
}



# ---------------------------------------------------------------------------
# Divider helper
# ---------------------------------------------------------------------------
def _divider() -> QFrame:
    line = QFrame()
    line.setFrameShape(QFrame.Shape.HLine)
    line.setStyleSheet("color: #1e2235; background: #1e2235;")
    line.setFixedHeight(1)
    return line


# ---------------------------------------------------------------------------
# Main window
# ---------------------------------------------------------------------------

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self._report: CheckReport | None = None
        self._worker: CheckWorker | None = None
        self._setup_window()
        self._build_ui()
        self._apply_palette()

    # ------------------------------------------------------------------
    # Window setup
    # ------------------------------------------------------------------

    def _setup_window(self):
        self.setWindowTitle("HWID 9527")
        self.setMinimumSize(1100, 720)
        self.resize(1280, 800)

    def _apply_palette(self):
        palette = QPalette()
        palette.setColor(QPalette.ColorRole.Window,          QColor("#0f1117"))
        palette.setColor(QPalette.ColorRole.WindowText,      QColor("#e2e8f0"))
        palette.setColor(QPalette.ColorRole.Base,            QColor("#13162a"))
        palette.setColor(QPalette.ColorRole.AlternateBase,   QColor("#161b2e"))
        palette.setColor(QPalette.ColorRole.Text,            QColor("#cbd5e1"))
        palette.setColor(QPalette.ColorRole.Button,          QColor("#1e2235"))
        palette.setColor(QPalette.ColorRole.ButtonText,      QColor("#e2e8f0"))
        palette.setColor(QPalette.ColorRole.Highlight,       QColor("#4f46e5"))
        palette.setColor(QPalette.ColorRole.HighlightedText, QColor("#ffffff"))
        self.setPalette(palette)

    # ------------------------------------------------------------------
    # UI construction
    # ------------------------------------------------------------------

    def _build_ui(self):
        root = QWidget()
        root.setStyleSheet("background-color: #0f1117;")
        self.setCentralWidget(root)

        vbox = QVBoxLayout(root)
        vbox.setSpacing(0)
        vbox.setContentsMargins(0, 0, 0, 0)

        vbox.addWidget(self._build_header())
        vbox.addWidget(_divider())

        # Scrollable body
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background: #0f1117; }")

        body = QWidget()
        body.setStyleSheet("background: #0f1117;")
        body_layout = QVBoxLayout(body)
        body_layout.setContentsMargins(24, 20, 24, 20)
        body_layout.setSpacing(16)

        # File upload
        upload_title = QLabel("1 — Upload Files")
        upload_title.setStyleSheet("color: #64748b; font-size: 11px; font-weight: 600;")
        body_layout.addWidget(upload_title)

        self._upload_panel = FileUploadPanel()
        self._upload_panel.file_changed.connect(self._on_file_changed)
        body_layout.addWidget(self._upload_panel)

        body_layout.addWidget(_divider())

        # Progress bar (hidden until running)
        self._progress_bar = QProgressBar()
        self._progress_bar.setRange(0, 100)
        self._progress_bar.setFixedHeight(6)
        self._progress_bar.setTextVisible(False)
        self._progress_bar.setStyleSheet("""
            QProgressBar {
                background: #1e2235; border-radius: 3px; border: none;
            }
            QProgressBar::chunk {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 #6366f1, stop:1 #06b6d4);
                border-radius: 3px;
            }
        """)
        self._progress_bar.setVisible(False)
        body_layout.addWidget(self._progress_bar)

        self._status_lbl = QLabel("")
        self._status_lbl.setStyleSheet("color: #64748b; font-size: 10px;")
        self._status_lbl.setVisible(False)
        body_layout.addWidget(self._status_lbl)

        # Results area title + summary
        results_header = QHBoxLayout()
        results_title = QLabel("2 — Check Results")
        results_title.setStyleSheet("color: #64748b; font-size: 11px; font-weight: 600;")
        results_header.addWidget(results_title)
        results_header.addStretch()
        self._region_summary_bar = SummaryBar()
        results_header.addWidget(self._region_summary_bar)
        body_layout.addLayout(results_header)

        # Tab widget
        self._tabs = QTabWidget()
        self._tabs.setStyleSheet("""
            QTabWidget::pane {
                border: 1px solid #1e2235;
                background: #0f1117;
                border-radius: 8px;
            }
            QTabBar::tab {
                background: #13162a;
                color: #64748b;
                padding: 8px 20px;
                font-size: 12px;
                border-top-left-radius: 8px;
                border-top-right-radius: 8px;
                margin-right: 2px;
            }
            QTabBar::tab:selected {
                background: #1e2235;
                color: #a5b4fc;
                font-weight: bold;
                border-bottom: 2px solid #6366f1;
            }
            QTabBar::tab:hover:!selected { background: #181c30; color: #cbd5e1; }
        """)

        # Region tab
        region_widget = QWidget()
        region_layout = QVBoxLayout(region_widget)
        region_layout.setContentsMargins(12, 12, 12, 12)
        self._region_summary = SummaryBar()
        region_layout.addWidget(self._region_summary)
        self._region_table = RegionResultTable()
        region_layout.addWidget(self._region_table)
        self._tabs.addTab(region_widget, "🌍  Region Check")

        # Device tab
        device_widget = QWidget()
        device_layout = QVBoxLayout(device_widget)
        device_layout.setContentsMargins(12, 12, 12, 12)
        self._device_summary = SummaryBar()
        device_layout.addWidget(self._device_summary)
        self._device_table = DeviceResultTable()
        device_layout.addWidget(self._device_table)
        self._tabs.addTab(device_widget, "🔧  Device Check")

        body_layout.addWidget(self._tabs, stretch=1)

        scroll.setWidget(body)
        vbox.addWidget(scroll, stretch=1)

    def _build_header(self) -> QWidget:
        header = QWidget()
        header.setFixedHeight(64)
        header.setStyleSheet(
            "background: qlineargradient("
            "x1:0, y1:0, x2:1, y2:0, "
            "stop:0 #0f1117, stop:0.5 #161b2e, stop:1 #0f1117);"
        )
        layout = QHBoxLayout(header)
        layout.setContentsMargins(24, 0, 24, 0)
        layout.setSpacing(12)

        # ---- LCFC brand block (red background, white text) ----
        lcfc_badge = QLabel("LCFC")
        lcfc_badge.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lcfc_badge.setFixedSize(54, 28)
        lcfc_badge.setStyleSheet("""
            background-color: #e50012;
            color: #ffffff;
            font-size: 16px;
            font-weight: 900;
            border-radius: 4px;
            font-family: 'Arial Black', Arial, sans-serif;
        """)

        powered_lbl = QLabel("Powered by SWDE-SIE")
        powered_lbl.setStyleSheet(
            "color: #94a3b8; font-size: 9px; background: transparent;"
        )

        brand_col = QVBoxLayout()
        brand_col.setSpacing(2)
        brand_col.setContentsMargins(0, 0, 0, 0)
        brand_col.setAlignment(Qt.AlignmentFlag.AlignVCenter)
        brand_col.addWidget(lcfc_badge, alignment=Qt.AlignmentFlag.AlignLeft)
        brand_col.addWidget(powered_lbl, alignment=Qt.AlignmentFlag.AlignLeft)

        layout.addLayout(brand_col)

        # Vertical separator
        sep = QFrame()
        sep.setFrameShape(QFrame.Shape.VLine)
        sep.setFixedWidth(1)
        sep.setStyleSheet("color: #1e2235; background: #1e2235;")
        layout.addWidget(sep)

        # Logo + title
        logo = QLabel("🔍")
        logo.setStyleSheet("font-size: 24px; background: transparent;")
        title = QLabel("HWID 9527")
        title.setStyleSheet(
            "font-size: 20px; font-weight: 800; color: #e2e8f0; background: transparent;"
        )
        subtitle = QLabel("Chromebook HWID Database Validator")
        subtitle.setStyleSheet(
            "font-size: 11px; color: #64748b; background: transparent;"
        )

        title_col = QVBoxLayout()
        title_col.setSpacing(2)
        title_col.addWidget(title)
        title_col.addWidget(subtitle)

        layout.addWidget(logo)
        layout.addLayout(title_col)
        layout.addStretch()

        # Project name label (populated after check)
        self._project_lbl = QLabel("")
        self._project_lbl.setStyleSheet(
            "color: #a5b4fc; font-size: 11px; background: transparent;"
        )
        layout.addWidget(self._project_lbl)

        # Run button
        self._run_btn = QPushButton("▶  Run Check")
        self._run_btn.setFixedSize(130, 36)
        self._run_btn.setEnabled(False)
        self._run_btn.clicked.connect(self._run_check)
        self._run_btn.setStyleSheet(self._run_btn_style(False))

        # Export button
        self._export_btn = QPushButton("⬇  Export")
        self._export_btn.setFixedSize(100, 36)
        self._export_btn.setEnabled(False)
        self._export_btn.clicked.connect(self._export)
        self._export_btn.setStyleSheet("""
            QPushButton {
                background: #1e2235; color: #94a3b8;
                border: 1px solid #2d3348; border-radius: 8px;
                font-size: 12px; font-weight: 600;
            }
            QPushButton:enabled { color: #e2e8f0; }
            QPushButton:enabled:hover { background: #252a40; border-color: #4f46e5; }
            QPushButton:disabled { color: #334155; }
        """)

        layout.addWidget(self._run_btn)
        layout.addWidget(self._export_btn)
        return header

    @staticmethod
    def _run_btn_style(enabled: bool) -> str:
        if enabled:
            return (
                "QPushButton {"
                "  background: qlineargradient(x1:0,y1:0,x2:1,y2:0,"
                "    stop:0 #6366f1, stop:1 #06b6d4);"
                "  color: #fff; border: none; border-radius: 8px;"
                "  font-size: 12px; font-weight: 700;"
                "}"
                "QPushButton:hover { background: qlineargradient(x1:0,y1:0,x2:1,y2:0,"
                "  stop:0 #4f46e5, stop:1 #0891b2); }"
                "QPushButton:pressed { background: #4338ca; }"
            )
        return (
            "QPushButton {"
            "  background: #1e2235; color: #334155;"
            "  border: 1px solid #2d3348; border-radius: 8px;"
            "  font-size: 12px; font-weight: 700;"
            "}"
        )

    # ------------------------------------------------------------------
    # Slots
    # ------------------------------------------------------------------

    def _on_file_changed(self):
        ready = self._upload_panel.is_ready()
        self._run_btn.setEnabled(ready)
        self._run_btn.setStyleSheet(self._run_btn_style(ready))

    def _run_check(self):
        if self._worker and self._worker.isRunning():
            return

        od, hwid, ec, avl = self._upload_panel.get_paths()

        self._run_btn.setEnabled(False)
        self._run_btn.setText("⏳  Checking…")
        self._run_btn.setStyleSheet(self._run_btn_style(False))
        self._export_btn.setEnabled(False)
        self._progress_bar.setValue(0)
        self._progress_bar.setVisible(True)
        self._status_lbl.setVisible(True)

        self._worker = CheckWorker(od, hwid, ec, avl, parent=self)
        self._worker.progress.connect(self._on_progress)
        self._worker.completed.connect(self._on_completed)
        self._worker.error.connect(self._on_error)
        self._worker.start()

    def _on_progress(self, message: str, percent: int):
        self._progress_bar.setValue(percent)
        self._status_lbl.setText(message)

    def _on_completed(self, report: CheckReport):
        self._report = report
        self._progress_bar.setValue(100)
        self._status_lbl.setText(
            f"✓ Check complete — {report.run_time}  |  "
            f"Project: {report.project_name}"
        )
        self._status_lbl.setStyleSheet("color: #22c55e; font-size: 10px;")
        self._project_lbl.setText(f"Project: {report.project_name}")

        # Populate tables
        self._region_table.load(report.region_results)
        self._device_table.load(report.device_results)

        # Update summaries
        r_sum = report.region_summary()
        d_sum = report.device_summary()
        self._region_summary.update_counts(r_sum)
        self._device_summary.update_counts(d_sum)

        # Re-enable buttons
        self._run_btn.setEnabled(True)
        self._run_btn.setText("▶  Run Check")
        self._run_btn.setStyleSheet(self._run_btn_style(True))
        self._export_btn.setEnabled(True)

    def _on_error(self, message: str):
        self._progress_bar.setVisible(False)
        self._status_lbl.setText("❌  Error occurred")
        self._status_lbl.setStyleSheet("color: #f87171; font-size: 10px;")
        self._run_btn.setEnabled(True)
        self._run_btn.setText("▶  Run Check")
        self._run_btn.setStyleSheet(self._run_btn_style(True))

        dlg = QMessageBox(self)
        dlg.setWindowTitle("Check Error")
        dlg.setIcon(QMessageBox.Icon.Critical)
        dlg.setText("An error occurred during the check.")
        dlg.setDetailedText(message)
        dlg.setStyleSheet("QMessageBox { background: #0f1117; color: #e2e8f0; }"
                          "QLabel { color: #e2e8f0; }")
        dlg.exec()

    # ------------------------------------------------------------------
    # Export
    # ------------------------------------------------------------------

    def _export(self):
        if not self._report:
            return

        default_name = (
            f"hwid_check_{self._report.project_name}_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Results", default_name,
            "Excel Files (*.xlsx)"
        )
        if not path:
            return

        try:
            self._write_excel(path)
            QMessageBox.information(
                self, "Export Complete",
                f"Results saved to:\n{path}"
            )
        except Exception as exc:
            QMessageBox.critical(
                self, "Export Error",
                f"Failed to save file:\n{exc}"
            )

    def _write_excel(self, path: str):
        wb = openpyxl.Workbook()

        # ---- Region sheet ----
        ws_r = wb.active
        ws_r.title = "Region Check"
        hdr_font = Font(bold=True, color="A5B4FC")
        hdr_fill = PatternFill("solid", fgColor="0F1030")
        region_headers = ["Geo", "Country", "Possible Codes", "Status", "Matched Code", "Note"]
        for ci, h in enumerate(region_headers, 1):
            cell = ws_r.cell(row=1, column=ci, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")

        for ri, r in enumerate(self._report.region_results, 2):
            row_data = [
                r.geo, r.country_name, ", ".join(r.possible_codes),
                r.status, r.matched_code, r.note,
            ]
            for ci, val in enumerate(row_data, 1):
                cell = ws_r.cell(row=ri, column=ci, value=val)
                if ci == 4:  # status column
                    cell.fill = _EXPORT_FILLS.get(r.status, PatternFill())
                    cell.font = _EXPORT_FONTS.get(r.status, Font())
                    cell.alignment = Alignment(horizontal="center")

        for col in ws_r.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws_r.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        # ---- Device sheet ----
        ws_d = wb.create_sheet("Device Check")
        device_headers = [
            "Category", "LCFC PN", "Vendor", "Description",
            "AVL Status", "HWID Key", "Check Status", "DB Status", "Note",
        ]
        for ci, h in enumerate(device_headers, 1):
            cell = ws_d.cell(row=1, column=ci, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")

        for ri, d in enumerate(self._report.device_results, 2):
            row_data = [
                d.category, d.lcfc_pn, d.vendor, d.description,
                d.avl_status, d.hwid_key, d.check_status,
                d.hwid_db_status, d.note,
            ]
            for ci, val in enumerate(row_data, 1):
                cell = ws_d.cell(row=ri, column=ci, value=val)
                if ci == 7:  # check status column
                    cell.fill = _EXPORT_FILLS.get(d.check_status, PatternFill())
                    cell.font = _EXPORT_FONTS.get(d.check_status, Font())
                    cell.alignment = Alignment(horizontal="center")

        for col in ws_d.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws_d.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

        wb.save(path)
